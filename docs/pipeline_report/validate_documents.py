"""Document-only structural and visual QA for the pipeline PDF/DOCX."""
from pathlib import Path
from datetime import datetime, timezone, timedelta
import hashlib
import json
import re
import sys
import zipfile
import xml.etree.ElementTree as ET

HERE=Path(__file__).resolve().parent
sys.path.insert(0,str(HERE.parent/'training_report'/'_build_deps'))
import pymupdf
from PIL import Image, ImageDraw, ImageFont
from docx import Document

STEM='MDD_Pipeline_hien_tai_Kien_truc_Model_Workflow'
qa=HERE/'qa';qa.mkdir(exist_ok=True)
source=json.loads((HERE/'pipeline_content.json').read_text(encoding='utf-8'))
references=json.loads((HERE.parent/'references_20261023.json').read_text(encoding='utf-8'))['references']
pdf=HERE/(STEM+'.pdf');word=HERE/(STEM+'.docx')
doc=pymupdf.open(pdf)
pages=[]
texts=[]
for i,page in enumerate(doc):
    text=page.get_text();texts.append(text)
    bad=[]
    for b in page.get_text('dict')['blocks']:
        x0,y0,x1,y1=b['bbox']
        if b['type']==0 and (x0<20 or x1>page.rect.width-20 or y0<5 or y1>page.rect.height-8):bad.append(b['bbox'])
    pix=page.get_pixmap(matrix=pymupdf.Matrix(1,1),alpha=False);pix.save(qa/f'page_{i+1:03d}.png')
    lines=[x.strip() for x in text.splitlines() if x.strip()]
    pages.append({'page':i+1,'characters':len(text),'suspect_bounds':bad,'first_lines':lines[4:7],'last_lines':lines[-3:]})
joined='\n'.join(texts);(qa/'pdf_text.txt').write_text(joined,encoding='utf-8')
primary_urls={row['url'] for row in references}
pdf_urls={link.get('uri') for page in doc for link in page.get_links() if link.get('uri')}
source_text=json.dumps(source,ensure_ascii=False)
cited=set(re.findall(r'\[(R-[A-Z0-9-]+)\]',source_text))
known={row['id'] for row in references}
assert cited <= known, f'Unknown research IDs: {cited-known}'
assert primary_urls <= pdf_urls, f'Missing PDF links: {primary_urls-pdf_urls}'
assert '(None)' not in joined and '\ue000' not in joined
font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',18)
for start in range(0,len(doc),9):
    sheet=Image.new('RGB',(1200,1750),'#dbe3eb');draw=ImageDraw.Draw(sheet)
    for slot,i in enumerate(range(start,min(start+9,len(doc)))):
        pic=Image.open(qa/f'page_{i+1:03d}.png').convert('RGB');pic.thumbnail((370,510))
        col,row=slot%3,slot//3
        sheet.paste(pic,(col*400+(400-pic.width)//2,row*580+40))
        draw.text((col*400+16,row*580+10),f'Trang {i+1}',font=font,fill='#172c43')
    sheet.save(qa/f'pdf_contact_{start//9+1:02d}.png')

d=Document(word)
word_urls={rel.target_ref for rel in d.part.rels.values()
           if rel.reltype.endswith('/hyperlink') and rel.is_external}
assert primary_urls <= word_urls, f'Missing DOCX links: {primary_urls-word_urls}'
heading_texts=[p.text for p in d.paragraphs if p.style.name=='Heading 1']
with zipfile.ZipFile(word) as z:
    xmls=[n for n in z.namelist() if n.endswith('.xml') or n.endswith('.rels')]
    for name in xmls:ET.fromstring(z.read(name))
    media=[n for n in z.namelist() if n.startswith('word/media/')]
    assert all(z.read(n).startswith(b'\x89PNG\r\n\x1a\n') for n in media)

result={'validated_at_bangkok':datetime.now(timezone(timedelta(hours=7))).isoformat(),
    'pdf':{'file':pdf.name,'sha256':hashlib.sha256(pdf.read_bytes()).hexdigest(),
        'pages':len(doc),'outline':doc.get_toc(),'replacement_characters':joined.count('\ufffd'),
        'caption_numbers':re.findall(r'(?m)^Hình (\d)\.',joined),'suspect_bounds':sum(len(p['suspect_bounds']) for p in pages),
        'pages_detail':pages},
    'docx':{'file':word.name,'sha256':hashlib.sha256(word.read_bytes()).hexdigest(),
        'xml_parts_parsed':len(xmls),'embedded_images':len(media),'inline_shapes':len(d.inline_shapes),
        'tables':len(d.tables),'headings_1':heading_texts,'paragraphs':len(d.paragraphs)},
    'research_references':{'count':len(references),'unknown_ids':sorted(cited-known),
        'pdf_primary_links':len(primary_urls & pdf_urls),'docx_primary_links':len(primary_urls & word_urls)},
    'scope':'PDF rendered and structurally inspected. DOCX ZIP/XML reopened and checked; no Word pagination/rendering engine used. No inference/training/TTS executed.'}
assert len(doc.get_toc())==len(source['chapters'])
assert joined.count('\ufffd')==0 and result['pdf']['suspect_bounds']==0
assert result['pdf']['caption_numbers']==['1','2','3','4']
assert len(d.inline_shapes)==4 and len(media)==4
assert all(ch['title'] in heading_texts for ch in source['chapters'])
(qa/'document_validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v if k!='pdf' else {a:b for a,b in v.items() if a!='pages_detail'} for k,v in result.items()},ensure_ascii=False,indent=2))
for p in pages:print(f"p{p['page']:02d} chars={p['characters']:4d} "+' / '.join(p['first_lines']))
