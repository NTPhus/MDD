"""Build matching PDF/DOCX architecture documents from a local JSON source."""
from pathlib import Path
from datetime import datetime, timezone, timedelta
import hashlib
import json
import re
import sys

HERE=Path(__file__).resolve().parent
TRAIN=HERE.parent/'training_report'
sys.path.insert(0,str(TRAIN/'_build_deps'))
sys.path.insert(0,str(TRAIN))
import build_report as base
import pymupdf
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas as pdfcanvas
from reportlab.platypus import Flowable, Paragraph, Preformatted, Spacer, PageBreak, NextPageTemplate, Table, TableStyle, KeepTogether
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus.tableofcontents import TableOfContents
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Mm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

WIDTH=base.WIDTH
STEM='MDD_Pipeline_hien_tai_Kien_truc_Model_Workflow'
pdfmetrics.registerFont(TTFont('Consolas','C:/Windows/Fonts/consola.ttf'))
CODE_STYLE=ParagraphStyle('PipelineCode',fontName='Consolas',fontSize=7.8,leading=11.5,
    backColor=colors.HexColor('#F0F3F7'),borderPadding=8,spaceBefore=2,spaceAfter=10)


class Diagram(Flowable):
    def __init__(self,spec):
        Flowable.__init__(self)
        self.spec=spec;self.width=WIDTH;self.height=spec['height']

    def draw(self):
        c=self.canv
        c.saveState()
        for edge in self.spec.get('edges',[]):
            points=edge['points']
            c.setStrokeColor(base.MUTED);c.setLineWidth(1)
            c.setDash(3,3) if edge.get('dashed') else c.setDash()
            path=c.beginPath();path.moveTo(*points[0])
            for point in points[1:]:path.lineTo(*point)
            c.drawPath(path)
            c.setDash()
            (x0,y0),(x,y)=points[-2:]
            dx,dy=x-x0,y-y0
            length=(dx*dx+dy*dy)**.5
            if length:
                ux,uy=dx/length,dy/length
                arrow=c.beginPath();arrow.moveTo(x,y)
                arrow.lineTo(x-6*ux+3*uy,y-6*uy-3*ux)
                arrow.lineTo(x-6*ux-3*uy,y-6*uy+3*ux);arrow.close()
                c.setFillColor(base.MUTED);c.drawPath(arrow,fill=1,stroke=0)
        for n in self.spec['nodes']:
            x,y,w,h=n['x'],n['y'],n['w'],n['h']
            kind=n.get('kind','normal')
            dark=kind=='dark'
            c.setFillColor(base.INK if dark else colors.HexColor('#FFF7E9') if kind=='legacy' else base.PALE if kind=='accent' else colors.HexColor('#F2F6F9'))
            c.setStrokeColor(base.INK if dark else colors.HexColor('#BA8D4B') if kind=='legacy' else base.TEAL if kind=='accent' else base.LINE)
            c.setDash(4,2) if kind=='legacy' else c.setDash()
            c.roundRect(x,y,w,h,5,fill=1,stroke=1);c.setDash()
            st=ParagraphStyle('Node',fontName='Arial',fontSize=n.get('font_size',8.6),leading=n.get('leading',12),
                textColor=colors.white if dark else base.INK,alignment=TA_CENTER)
            text='<b>'+base.clean(n['title'])+'</b>'
            if n.get('body'):text+='<br/>'+base.clean(n['body'])
            paragraph=Paragraph(text,st)
            _,ph=paragraph.wrap(w-16,h-10)
            if ph>h-10:raise ValueError(f"Diagram node overflow: {n['title']}")
            paragraph.drawOn(c,x+8,y+(h-ph)/2)
        c.restoreState()


def node(x,y,w,h,title,body='',kind='normal',**kw):return dict(x=x,y=y,w=w,h=h,title=title,body=body,kind=kind,**kw)
def edge(*points,dashed=False):return dict(points=points,dashed=dashed)
W=WIDTH
DIAGRAMS={
 'architecture':dict(height=402,nodes=[
    node(0,352,W,43,'Browser: demo/chat.js','Mic / upload · hội thoại · feedback · phát voice','dark'),
    node(0,291,W,43,'HTTP server: demo.py + ChatService','POST audio → job; GET polling → text/speech snapshot','accent'),
    node(0,230,W,43,'JSONL IPC: local_chat_process.py','Một worker riêng; event stage / speech / chunk / done'),
    node(0,16,W,195,'','',kind='normal'),
    node(12,192,W-24,17,'Worker: local_chat_worker.py',font_size=7,leading=7),
    node(12,132,223,57,'ASR chữ: wav2vec2-vi-en-v1','Audio → transcript + word intervals\nCPU'),
    node(12,54,223,62,'Phoneme: cascade-B','Audio → phone CTC posteriors\nCMU variants + assessment · CPU'),
    node(252,54,W-264,135,'Một Qwen3-0.6B base dùng chung','Base: normalize + dialogue\nLoRA grammar: 2 task viết lại\nLoRA feedback: có sẵn, nhánh legacy\nRuntime Qwen: auto / cpu / cuda','accent'),
 ],edges=[edge((W/2,352),(W/2,334)),edge((W/2,291),(W/2,273)),edge((W/2,230),(W/2,211))]),
 'workflow':dict(height=446,nodes=[
    node(13,392,W-26,45,'1. Nhận audio hợp lệ','PCM16 WAV mono 16 kHz, 0,25–30 giây → tạo job','dark'),
    node(13,337,W-26,45,'2. wav2vec2 ASR trên CPU','Sinh raw transcript và khoảng thời gian từ'),
    node(13,282,W-26,45,'3. wav2vec2 cascade-B + assessment trên CPU','Posterior + reference CMU → assessment → uncertain / gợi ý'),
    node(13,227,W-26,45,'4. Qwen base chuẩn hóa hình thức','Chỉ hội thoại tự do; lexical guard giữ nội dung đã nói'),
    node(13,172,W-26,45,'5. Qwen base sinh câu trả lời','Sampling hội thoại → lọc lặp / kiểm tra policy / retry tối đa 1'),
    node(13,117,W-26,45,'6. Công bố reply_ready','Câu trả lời cố định; UI/voice có thể bắt đầu trong khi feedback chờ','accent'),
    node(13,62,W-26,45,'7. Grammar feedback cho hội thoại tự do','Hai lượt greedy: grammar_corrected rồi natural; guard từng trường'),
    node(13,7,W-26,45,'8. Kết thúc job','Luồng assessment mới bỏ qua legacy coach → ready / cancelled / error'),
 ],edges=[edge((W/2,y),(W/2,y-10)) for y in [392,337,282,227,172,117,62]]),
 'routing':dict(height=316,nodes=[
    node(65,259,W-130,46,'Một backbone Qwen3-0.6B','Cache trong worker; tải model theo runtime policy','dark'),
    node(0,91,153,135,'qwen-base','disable_adapter() nếu PEFT đã nạp\n\nNormalize transcript\nSinh dialogue\nFallback rewrite','accent',font_size=8.3),
    node(170,91,153,135,'qwen-grammar','Bật LoRA grammar\n\nGrammar correction\nNatural wording\nHai generation riêng','accent',font_size=8.3),
    node(340,91,W-340,135,'qwen-trained','Bật LoRA cascade feedback\n\nCoach phát âm legacy\nKhông gọi trong speech path có assessment','legacy',font_size=8.3),
    node(0,6,W,62,'Chọn model ở UI không đổi backbone reply speech','Trong speech chat, dialogue luôn qwen-base. Metadata coaching_model_id không tự chứng minh adapter đã được gọi.','normal'),
 ],edges=[edge((W/2,259),(W/2,244),(76.5,244),(76.5,226)),edge((W/2,259),(W/2,244),(246.5,244),(246.5,226)),edge((W/2,259),(W/2,244),(416,244),(416,226),dashed=True)]),
 'reference':dict(height=352,nodes=[
    node(170,293,W-170,48,'POST /api/chat/reference','Cùng job_id; user xác nhận câu thực sự đã nói','dark'),
    node(0,204,150,65,'Posterior cache phía server','Audio evidence ban đầu\nKèm checkpoint/vocab hash','accent'),
    node(170,204,W-170,65,'Reassess theo text mới','Không decode audio, chạy ASR hoặc phoneme model lại\nSo spoken_text_key'),
    node(170,111,150,65,'Cùng nội dung từ','Cập nhật reference\nGiữ reply / reply voice\nKhông gọi Qwen'),
    node(340,111,W-340,65,'Nội dung từ thay đổi','revision++\nXóa reply/voice cũ\nGiữ audio evidence','accent'),
    node(170,12,W-170,63,'Regenerate bằng text đã xác nhận','Qwen base trả lời lại → grammar rewrite\nCùng lượt; không tạo bản thu mới','accent'),
 ],edges=[edge((330,293),(330,269)),edge((150,236),(170,236)),edge((330,204),(330,190),(245,190),(245,176)),edge((330,204),(330,190),(416,190),(416,176)),edge((416,111),(416,75))]),
}


class PipelineDoc(base.ReportDoc):
    def page_bg(self,c,doc):
        w,h=c._pagesize;c.saveState();c.setStrokeColor(base.LINE)
        c.line(base.MARGIN,h-15*mm,w-base.MARGIN,h-15*mm)
        c.setFont('Arial-Bold',7.5);c.setFillColor(base.TEAL)
        c.drawString(base.MARGIN,h-12*mm,'MDD  /  Pipeline hiện tại')
        c.setFont('Arial',7);c.setFillColor(base.MUTED)
        c.drawRightString(w-base.MARGIN,h-12*mm,'Kiến trúc · Model · Workflow  |  23/10/2026')
        c.line(base.MARGIN,14*mm,w-base.MARGIN,14*mm)
        c.drawString(base.MARGIN,10*mm,'Snapshot từ code và config thực tế; đường chat và practice được mô tả riêng')
        c.drawRightString(w-base.MARGIN,10*mm,str(doc.page));c.restoreState()


def export_diagrams():
    folder=HERE/'figures';folder.mkdir(exist_ok=True)
    manifest=[]
    for key,spec in DIAGRAMS.items():
        target=folder/(key+'.pdf')
        c=pdfcanvas.Canvas(str(target),pagesize=(WIDTH,spec['height']))
        Diagram(spec).drawOn(c,0,0);c.save()
        d=pymupdf.open(target);pix=d[0].get_pixmap(matrix=pymupdf.Matrix(2.5,2.5),alpha=False)
        png=folder/(key+'.png');pix.save(png)
        manifest.append({'id':key,'pdf':str(target.relative_to(HERE)),'png':str(png.relative_to(HERE)),
            'png_sha256':hashlib.sha256(png.read_bytes()).hexdigest()})
    (HERE/'diagrams_manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')


def build_pdf(source):
    out=HERE/(STEM+'.pdf')
    doc=PipelineDoc(str(out),title='MDD – Pipeline hiện tại: kiến trúc, model và workflow',author='Dự án MDD',subject='Tài liệu kiến trúc và vận hành hiện tại')
    label=ParagraphStyle('CoverLabel',fontName='Arial-Bold',fontSize=11,leading=17,textColor=base.TEAL,spaceAfter=17)
    title=ParagraphStyle('CoverTitle',fontName='Arial-Bold',fontSize=32,leading=42,textColor=base.INK,spaceAfter=20)
    sub=ParagraphStyle('CoverSub',fontName='Arial',fontSize=15,leading=23,textColor=base.MUTED,spaceAfter=18)
    story=[Spacer(1,25*mm),Paragraph('TÀI LIỆU KIẾN TRÚC & VẬN HÀNH',label),
        Paragraph('Pipeline hiện tại<br/>của hệ thống MDD',title),
        Paragraph('Cấu trúc hệ thống<br/>Model nào làm việc gì<br/>Workflow từ audio đến phản hồi',sub),
        Spacer(1,8*mm),base.para('**MDD – Phát hiện và chẩn đoán lỗi phát âm tiếng Anh**\nTài liệu đồng hành với báo cáo huấn luyện wav2vec2 và Qwen.'),Spacer(1,6*mm),
        base.table({'headers':['Phạm vi','Snapshot'], 'rows':[
            ['Trọng tâm','Chat giọng nói hiện tại; thêm đường luyện phát âm và cấu trúc dự án'],
            ['Nguồn','Code, model registry, artifact readiness và GET /api/chat/config'],
            ['Thời điểm',source['snapshot_time']],
            ['Định dạng','PDF để đọc/chia sẻ; DOCX có cùng nội dung để chỉnh sửa'],
        ],'widths':[1,3]}),Spacer(1,9*mm),base.para('Sơ đồ phản ánh đường chạy trong source đang có. Nhánh nghiên cứu, nhánh legacy và model đã train nhưng chưa gọi trong luồng chính được ghi rõ.','small'),
        NextPageTemplate('body'),PageBreak(),base.para('Mục lục','h1')]
    toc=TableOfContents();toc.levelStyles=[ParagraphStyle('TOC',fontName='Arial-Bold',fontSize=10,leading=15,spaceBefore=8,textColor=base.INK)]
    story += [toc,Spacer(1,12*mm),base.para('Cách đọc tài liệu','h2'),
        base.para('“Đang dùng” chỉ đường gọi có trong workflow hiện tại. “Ready” cho biết artifact/config đủ điều kiện nạp; không đồng nghĩa một lượt xử lý đã thực sự gọi model đó. Device=auto là chính sách chọn thiết bị, khác device thực tế được ghi cho mỗi lần sinh.'),
        base.para('Các sơ đồ kiến trúc được dựng từ code để giải thích pipeline, không phải đồ thị kết quả huấn luyện. Tham chiếu [Pxx] chỉ đến danh mục file/hàm ở cuối tài liệu.')]
    for i,ch in enumerate(source['chapters'],1):
        story.append(PageBreak());heading=base.para(ch['title'],'h1');heading.report_heading=(0,ch['title'],f'pipeline-{i}');story.append(heading)
        for b in ch['blocks']:
            kind=b['type']
            if kind=='code':story.append(Preformatted(b['text'],CODE_STYLE,maxLineLength=100,splitChars=' /'))
            elif kind in ['p','small','source','h2','h3']:story.append(base.para(b['text'],{'p':'body'}.get(kind,kind)))
            elif kind=='bullets':story.extend(base.para('• '+x,'bullet') for x in b['items']);story.append(Spacer(1,4))
            elif kind=='table':
                story += [base.table(b),Spacer(1,8)]
                if b.get('source'):story.append(base.para(b['source'],'source'))
            elif kind=='note':
                t=Table([[base.para(b['text'])]],colWidths=[WIDTH]);t.setStyle(TableStyle([
                    ('BACKGROUND',(0,0),(-1,-1),base.PALE),('BOX',(0,0),(-1,-1),.6,base.TEAL),
                    ('LEFTPADDING',(0,0),(-1,-1),10),('RIGHTPADDING',(0,0),(-1,-1),10),
                    ('TOPPADDING',(0,0),(-1,-1),8),('BOTTOMPADDING',(0,0),(-1,-1),3)]));story += [t,Spacer(1,9)]
            elif kind=='diagram':story += [KeepTogether([Diagram(DIAGRAMS[b['id']]),base.para(b['caption'],'caption'),base.para(b.get('source',''),'source')]),Spacer(1,8)]
            elif kind=='pagebreak':story.append(PageBreak())
            else:raise ValueError(kind)
    doc.multiBuild(story)
    return out


def add_rich(paragraph,text):
    for part in re.split(r'(\*\*.*?\*\*|`[^`]+`)',str(text)):
        if part.startswith('**') and part.endswith('**'):paragraph.add_run(part[2:-2]).bold=True
        elif part.startswith('`') and part.endswith('`'):
            run=paragraph.add_run(part[1:-1]);run.font.color.rgb=RGBColor.from_string('087E8B');run.font.size=Pt(9)
        else:paragraph.add_run(part)


def word_table(doc,b):
    t=doc.add_table(rows=1, cols=len(b['headers']));t.style='Light Shading Accent 1'
    t.autofit=False
    ratios=b.get('widths',[1]*len(b['headers']))
    for i,label in enumerate(b['headers']):
        cell=t.rows[0].cells[i];cell.width=Mm(174*ratios[i]/sum(ratios));add_rich(cell.paragraphs[0],label)
        for r in cell.paragraphs[0].runs:r.bold=True
    header=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(header)
    for row in b['rows']:
        cells=t.add_row().cells
        for i,value in enumerate(row):
            cells[i].width=Mm(174*ratios[i]/sum(ratios));add_rich(cells[i].paragraphs[0],'—' if value is None else str(value))
    for row in t.rows:
        cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(3)
                for r in p.runs:r.font.size=Pt(8.2)
    if b.get('source'):
        p=doc.add_paragraph();add_rich(p,b['source']);p.style='Caption'


def build_docx(source):
    doc=Document();section=doc.sections[0]
    section.page_width=Mm(210);section.page_height=Mm(297)
    section.top_margin=Mm(19);section.bottom_margin=Mm(18);section.left_margin=section.right_margin=Mm(18)
    normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10)
    normal.paragraph_format.line_spacing=1.25;normal.paragraph_format.space_after=Pt(7)
    for name,size in [('Title',30),('Heading 1',20),('Heading 2',13),('Heading 3',11)]:
        st=doc.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('172C43' if name!='Heading 2' else '087E8B')
    doc.styles['Caption'].font.name='Arial';doc.styles['Caption'].font.size=Pt(8)
    header=section.header.paragraphs[0];header.text='MDD  |  Pipeline hiện tại · Model · Workflow';header.style='Caption'
    foot=section.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    foot.add_run('MDD · 23/10/2026   |   ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
    for r in foot.runs:r.font.size=Pt(8)
    doc.add_paragraph('TÀI LIỆU KIẾN TRÚC & VẬN HÀNH','Subtitle')
    doc.add_paragraph('Pipeline hiện tại\ncủa hệ thống MDD','Title')
    doc.add_paragraph('Cấu trúc hệ thống · Model nào làm việc gì · Workflow từ audio đến phản hồi','Subtitle')
    doc.add_paragraph('Tài liệu đồng hành với báo cáo huấn luyện wav2vec2 và Qwen.\nSnapshot: '+source['snapshot_time'])
    word_table(doc,{'headers':['Phạm vi','Nội dung'],'rows':[
        ['Trọng tâm','Chat giọng nói, luyện phát âm, runtime và cấu trúc code hiện tại'],
        ['Nguồn','Code + artifact metadata + GET config, không chạy inference/train/TTS'],
        ['Sơ đồ','Kiến trúc, thứ tự xử lý, Qwen adapter routing và reference revision'],
    ],'widths':[1,3]})
    doc.add_page_break();doc.add_heading('Nội dung',1)
    for ch in source['chapters']:doc.add_paragraph(ch['title'])
    doc.add_heading('Cách đọc',2)
    doc.add_paragraph('Ready là trạng thái artifact, khác việc model thực sự được gọi. Device=auto là policy; metadata từng generation mới cho biết thiết bị thực tế. Tài liệu phân biệt runtime hiện tại, nhánh legacy và nhánh nghiên cứu.')
    for ch in source['chapters']:
        doc.add_page_break();doc.add_heading(ch['title'],1)
        for b in ch['blocks']:
            kind=b['type']
            if kind in ['h2','h3']:doc.add_heading(b['text'],2 if kind=='h2' else 3)
            elif kind in ['p','small','source','code','note']:
                para=doc.add_paragraph();add_rich(para,b['text'])
                if kind in ['small','source']:para.style='Caption'
                if kind=='note':
                    shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'EAF4F5');para._p.get_or_add_pPr().append(shading)
                if kind=='code':
                    for r in para.runs:r.font.name='Consolas';r.font.size=Pt(8)
            elif kind=='bullets':
                for item in b['items']:add_rich(doc.add_paragraph(style='List Bullet'),item)
            elif kind=='table':word_table(doc,b)
            elif kind=='diagram':
                para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.keep_with_next=True
                para.add_run().add_picture(str(HERE/'figures'/(b['id']+'.png')),width=Mm(174))
                cap=doc.add_paragraph(style='Caption');add_rich(cap,b['caption'])
                cap.paragraph_format.keep_with_next=bool(b.get('source'))
                if b.get('source'):add_rich(doc.add_paragraph(style='Caption'),b['source'])
            elif kind=='pagebreak':doc.add_page_break()
    out=HERE/(STEM+'.docx');doc.save(out);return out


def main():
    source=json.loads((HERE/'pipeline_content.json').read_text(encoding='utf-8'))
    export_diagrams();pdf=build_pdf(source);docx=build_docx(source)
    manifest={'built_at_bangkok':datetime.now(timezone(timedelta(hours=7))).isoformat(),
        'source':'pipeline_content.json','chapters':len(source['chapters']),'diagrams':list(DIAGRAMS),
        'outputs':[{'file':p.name,'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in [pdf,docx]]}
    (HERE/'build_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(manifest,ensure_ascii=False,indent=2))

if __name__=='__main__':main()
